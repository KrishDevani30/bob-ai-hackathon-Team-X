"""CAPA report export to DOCX and PDF."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from typing import Any

from backend.capa.generator import CapaReport


def export_docx(report: CapaReport, output_path: Path) -> Path:
    """Export a CapaReport to a Word (.docx) document."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise ImportError("python-docx not installed. Run: pip install python-docx") from exc

    doc = Document()

    # Title
    title = doc.add_heading("CORRECTIVE AND PREVENTIVE ACTION REPORT", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Draft watermark paragraph
    draft_para = doc.add_paragraph(
        "[DRAFT] System-generated. Requires qualified human review before regulatory use."
    )
    draft_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    draft_run = draft_para.runs[0]
    draft_run.bold = True
    draft_run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    doc.add_paragraph()

    # Header table
    table = doc.add_table(rows=8, cols=2)
    table.style = "Table Grid"
    header_data = [
        ("Study ID", report.study_id),
        ("Protocol Version", report.protocol_version),
        ("Site ID", report.site_id),
        ("Principal Investigator", report.pi_name),
        ("Report Date", report.report_date),
        ("Risk Band", report.risk_band),
        ("Risk Score", f"{report.risk_score:.1f} / 100"),
        ("Total Deviations", str(report.total_deviations)),
    ]
    for i, (label, value) in enumerate(header_data):
        row = table.rows[i]
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = value

    doc.add_paragraph()
    doc.add_heading("Deviation Summary", level=2)

    # Deviation summary table
    if report.deviation_summary:
        cols = ["Deviation Type", "Count", "Affected Patients", "Severity Distribution"]
        tbl = doc.add_table(rows=1 + len(report.deviation_summary), cols=len(cols))
        tbl.style = "Table Grid"
        for j, col in enumerate(cols):
            cell = tbl.rows[0].cells[j]
            cell.text = col
            cell.paragraphs[0].runs[0].bold = True
        for i, row_data in enumerate(report.deviation_summary, start=1):
            row = tbl.rows[i]
            sev_dist = "; ".join(
                f"{k}: {v}" for k, v in row_data.get("severity_distribution", {}).items()
            )
            row.cells[0].text = row_data.get("deviation_type", "")
            row.cells[1].text = str(row_data.get("count", ""))
            row.cells[2].text = str(row_data.get("affected_patients", ""))
            row.cells[3].text = sev_dist

    # Narrative sections
    sections = [
        ("Root Cause Analysis", report.capa_content.root_cause_analysis),
        ("Immediate Corrective Action", report.capa_content.immediate_corrective_action),
        ("Preventive Action", report.capa_content.preventive_action),
        ("Effectiveness Check Criteria", report.capa_content.effectiveness_check_criteria),
    ]
    for heading, content in sections:
        doc.add_paragraph()
        doc.add_heading(heading, level=2)
        doc.add_paragraph(content)

    # Regulatory references
    doc.add_paragraph()
    doc.add_heading("Regulatory References", level=2)
    for ref in report.capa_content.regulatory_references:
        doc.add_paragraph(f"* {ref}")

    # Signature block
    doc.add_paragraph()
    doc.add_heading("Review and Approval", level=2)
    sig_table = doc.add_table(rows=3, cols=3)
    sig_table.style = "Table Grid"
    for j, label in enumerate(["Name", "Role", "Signature / Date"]):
        cell = sig_table.rows[0].cells[j]
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
    for row in sig_table.rows[1:]:
        for cell in row.cells:
            cell.text = " "

    # Footer via section properties
    section = doc.sections[0]
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.text = (
        f"DRAFT | {report.generated_by} | Generated: {date.today().isoformat()} | "
        f"Study: {report.study_id} | Site: {report.site_id}"
    )
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(output_path))
    return output_path


def export_pdf(report: CapaReport, output_path: Path) -> Path:
    """Export a CapaReport to PDF using reportlab."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        )
    except ImportError as exc:
        raise ImportError("reportlab not installed. Run: pip install reportlab") from exc

    styles = getSampleStyleSheet()
    warning_style = ParagraphStyle(
        "Warning",
        parent=styles["Normal"],
        textColor=colors.red,
        fontName="Helvetica-Bold",
        fontSize=10,
        spaceAfter=6,
    )
    h1_style = ParagraphStyle(
        "H1", parent=styles["Heading1"], fontSize=14, spaceAfter=8
    )
    h2_style = ParagraphStyle(
        "H2", parent=styles["Heading2"], fontSize=11, spaceAfter=6
    )
    body_style = styles["Normal"]

    def _page_footer(canvas, doc):  # type: ignore[no-untyped-def]
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        footer_text = (
            f"DRAFT | {report.generated_by} | "
            f"Study: {report.study_id} | Site: {report.site_id} | "
            f"Page {doc.page}"
        )
        canvas.drawCentredString(A4[0] / 2, 1.5 * cm, footer_text)
        canvas.restoreState()

    story = []

    story.append(Paragraph("CORRECTIVE AND PREVENTIVE ACTION REPORT", h1_style))
    story.append(Paragraph(
        "[DRAFT] System-generated. Requires qualified human review before regulatory use.",
        warning_style,
    ))
    story.append(Spacer(1, 0.4 * cm))

    # Header table
    header_data = [
        ["Study ID", report.study_id, "Risk Band", report.risk_band],
        ["Protocol Version", report.protocol_version, "Risk Score", f"{report.risk_score:.1f} / 100"],
        ["Site ID", report.site_id, "Total Deviations", str(report.total_deviations)],
        ["PI", report.pi_name, "Report Date", report.report_date],
    ]
    header_tbl = Table(header_data, colWidths=[3.5 * cm, 6 * cm, 3.5 * cm, 4 * cm])
    header_tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (0, -1), colors.lightgrey),
        ("BACKGROUND", (2, 0), (2, -1), colors.lightgrey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(header_tbl)
    story.append(Spacer(1, 0.5 * cm))

    # Deviation summary
    story.append(Paragraph("Deviation Summary", h2_style))
    if report.deviation_summary:
        tbl_data = [["Deviation Type", "Count", "Affected Patients", "Severity Distribution"]]
        for row in report.deviation_summary:
            sev = "; ".join(f"{k}: {v}" for k, v in row.get("severity_distribution", {}).items())
            tbl_data.append([
                row.get("deviation_type", ""),
                str(row.get("count", "")),
                str(row.get("affected_patients", "")),
                sev,
            ])
        dev_tbl = Table(tbl_data, colWidths=[4 * cm, 2 * cm, 3 * cm, 8 * cm])
        dev_tbl.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3c5e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(dev_tbl)
    story.append(Spacer(1, 0.4 * cm))

    # Narrative sections
    sections = [
        ("Root Cause Analysis", report.capa_content.root_cause_analysis),
        ("Immediate Corrective Action", report.capa_content.immediate_corrective_action),
        ("Preventive Action", report.capa_content.preventive_action),
        ("Effectiveness Check Criteria", report.capa_content.effectiveness_check_criteria),
    ]
    for heading, content in sections:
        story.append(Paragraph(heading, h2_style))
        story.append(Paragraph(content, body_style))
        story.append(Spacer(1, 0.3 * cm))

    # Regulatory references
    story.append(Paragraph("Regulatory References", h2_style))
    for ref in report.capa_content.regulatory_references:
        story.append(Paragraph(f"* {ref}", body_style))

    story.append(Spacer(1, 0.5 * cm))
    story.append(Paragraph("Review and Approval", h2_style))
    sig_data = [
        ["Name", "Role / Title", "Signature", "Date"],
        [" ", " ", " ", " "],
        [" ", " ", " ", " "],
    ]
    sig_tbl = Table(sig_data, colWidths=[4.5 * cm, 4.5 * cm, 4.5 * cm, 3.5 * cm])
    sig_tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("ROWHEIGHT", (0, 1), (-1, -1), 1.5 * cm),
    ]))
    story.append(sig_tbl)

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2.5 * cm,
    )
    doc.build(story, onFirstPage=_page_footer, onLaterPages=_page_footer)
    return output_path
